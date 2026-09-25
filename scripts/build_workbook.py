"""Account CoPilot — REPORTER: builds the Master Book workbook.

Usage: python3 scripts/build_workbook.py [output.xlsx]
"""
import os
import sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule
from openpyxl.chart import BarChart, PieChart, Reference
from openpyxl.chart.label import DataLabelList

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from load_data import load_all, ROOT, RUN_DATE, SIGNAL_ORDER, name_key  # noqa: E402

FONT = "Arial"
# Carbon-fibre blue palette
CARBON_FG, CARBON_BG = "1B263B", "0B1320"
HEADER = "15325A"
HEADER_2 = "1F4E79"
BAND = "EEF3FA"
INPUT_FILL = "FFF3D1"
INPUT_BORDER = "2E75B6"
ACCENT = "4FA3F7"
TAB_COLORS = ["0B1320", "0F1D33", "14284A", "15325A", "1A3D6B", "1F4E79", "245C8F",
              "2A6AA5", "2E75B6", "3A86C8", "4A96D8", "5AA6E6", "6BB5F0"]
SIGNAL_FILLS = {"VERY STRONG SIGNAL": ("0B6E4F", "FFFFFF"), "STRONG SIGNAL": ("2E9E6A", "FFFFFF"),
                "MODERATE SIGNAL": ("F2C14E", "1B263B"), "WEAK SIGNAL": ("F7E7B4", "1B263B"),
                "NO SIGNAL": ("D9DEE5", "1B263B"), "CONFLICTING SIGNAL": ("E06C75", "FFFFFF")}

thin = Side(style="thin", color="C9D3E0")
frame = Side(style="medium", color=INPUT_BORDER)
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
INPUT_BORDER_STYLE = Border(left=frame, right=frame, top=frame, bottom=frame)
FIRST_ROW = 4  # header row is 3
MAX_ROW = 2000

wb = Workbook()
COLMAP = {}  # sheet -> key -> column letter


def banner(ws, title, subtitle, ncols):
    last = get_column_letter(max(ncols, 6))
    ws.merge_cells(f"A1:{last}1")
    ws.merge_cells(f"A2:{last}2")
    ws["A1"] = title
    ws["A2"] = subtitle
    carbon = PatternFill(patternType="darkTrellis", fgColor=CARBON_FG, bgColor=CARBON_BG)
    for col in range(1, max(ncols, 6) + 1):
        ws.cell(row=1, column=col).fill = carbon
        ws.cell(row=2, column=col).fill = PatternFill("solid", fgColor=CARBON_BG)
    ws["A1"].font = Font(name=FONT, size=16, bold=True, color="FFFFFF")
    ws["A2"].font = Font(name=FONT, size=9, italic=True, color="A9C7EC")
    ws["A1"].alignment = Alignment(vertical="center", indent=1)
    ws["A2"].alignment = Alignment(vertical="center", indent=1)
    ws.row_dimensions[1].height = 30
    ws.row_dimensions[2].height = 18


def write_table(ws, name, title, subtitle, cols, rows, input_keys=(), tab=0, validations=None):
    """cols: list of (header, key, width, kind) kind in {'text','url','num','wrap'}"""
    ws.sheet_properties.tabColor = TAB_COLORS[tab % len(TAB_COLORS)]
    banner(ws, title, subtitle, len(cols))
    COLMAP[ws.title] = {}
    for ci, (hdr, key, width, kind) in enumerate(cols, 1):
        L = get_column_letter(ci)
        COLMAP[ws.title][key] = L
        c = ws.cell(row=3, column=ci, value=hdr)
        is_input = key in input_keys
        c.fill = PatternFill("solid", fgColor="B8860B" if is_input else HEADER)
        c.font = Font(name=FONT, bold=True, color="FFFFFF", size=10)
        c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        c.border = BORDER
        ws.column_dimensions[L].width = width
    ws.row_dimensions[3].height = 34
    for ri, r in enumerate(rows, FIRST_ROW):
        band = PatternFill("solid", fgColor=BAND) if ri % 2 == 0 else PatternFill("solid", fgColor="FFFFFF")
        for ci, (hdr, key, width, kind) in enumerate(cols, 1):
            v = r.get(key, "") if isinstance(r, dict) else r[ci - 1]
            if isinstance(v, (list, dict)):
                v = "; ".join(str(x) for x in v) if isinstance(v, list) else str(v)
            if v is None:
                v = ""
            c = ws.cell(row=ri, column=ci, value=v)
            c.font = Font(name=FONT, size=9)
            c.border = BORDER
            c.alignment = Alignment(wrap_text=kind in ("wrap",), vertical="top")
            if key in input_keys:
                c.fill = PatternFill("solid", fgColor=INPUT_FILL)
                c.border = INPUT_BORDER_STYLE
            else:
                c.fill = band
            if kind == "url" and isinstance(v, str) and v.startswith("http"):
                c.hyperlink = v
                c.font = Font(name=FONT, size=9, color="1F5FBF", underline="single")
            if kind == "num" and v != "":
                c.number_format = "#,##0"
    last_row = max(FIRST_ROW, FIRST_ROW + len(rows) - 1)
    ref = f"A3:{get_column_letter(len(cols))}{last_row}"
    if rows:
        t = Table(displayName=name, ref=ref)
        t.tableStyleInfo = TableStyleInfo(name="TableStyleLight9", showRowStripes=False)
        ws.add_table(t)
    ws.freeze_panes = ws.cell(row=FIRST_ROW, column=3)
    # signal conditional formatting
    for key in ("s2p_signal_level", "account_s2p_signal", "level"):
        if key in COLMAP[ws.title]:
            L = COLMAP[ws.title][key]
            rng = f"{L}{FIRST_ROW}:{L}{MAX_ROW}"
            for sig, (bg, fg) in SIGNAL_FILLS.items():
                ws.conditional_formatting.add(rng, FormulaRule(
                    formula=[f'${L}{FIRST_ROW}="{sig}"'],
                    fill=PatternFill("solid", fgColor=bg, bgColor=bg),
                    font=Font(name=FONT, bold=True, color=fg)))
    for key, options in (validations or {}).items():
        if key in COLMAP[ws.title]:
            L = COLMAP[ws.title][key]
            dv = DataValidation(type="list", formula1='"' + ",".join(options) + '"', allow_blank=True)
            ws.add_data_validation(dv)
            dv.add(f"{L}{FIRST_ROW}:{L}{MAX_ROW}")
    return last_row


def col(sheet, key):
    L = COLMAP[sheet][key]
    return f"'{sheet}'!${L}${FIRST_ROW}:${L}${MAX_ROW}"


def main(out):
    d = load_all()
    A, P = d["accounts"], d["contacts"]
    for a in A:
        a["revenue_usd_m"] = a.get("revenue_usd_m") or ""
        a["third_party_apps_text"] = "; ".join(
            f"{x.get('name')} ({x.get('category','')}, {x.get('status','')})" for x in a.get("third_party_apps") or [])
    by_co = {}
    for p in P:
        by_co.setdefault(p["company_id"], []).append(p)

    ws_dash = wb.active
    ws_dash.title = "Executive Dashboard"

    # ---------------- Accounts ----------------
    ws = wb.create_sheet("Accounts")
    acc_cols = [
        ("Company ID", "company_id", 10, "text"), ("Company Name", "company_name", 30, "text"),
        ("Company Website", "company_website", 26, "url"), ("Country", "country", 9, "text"),
        ("Exchange", "exchange", 10, "text"), ("Ticker", "ticker", 10, "text"),
        ("Industry", "industry", 20, "text"), ("ICP Fit", "icp_fit", 10, "text"),
        ("ICP Fit Reason", "icp_fit_reason", 34, "wrap"),
        ("Estimated Revenue (USD m)", "revenue_usd_m", 13, "num"), ("Revenue (Local)", "revenue_local", 16, "text"),
        ("Revenue FY", "revenue_fy", 10, "text"), ("Revenue Source", "revenue_source_url", 24, "url"),
        ("Employee Range", "employee_range", 14, "text"), ("Ownership", "ownership", 30, "wrap"),
        ("Parent Company", "parent_company", 22, "text"), ("Subsidiary", "subsidiaries", 34, "wrap"),
        ("Board Phone", "board_phone", 16, "text"),
        ("Procurement Model", "procurement_model", 30, "wrap"), ("ERP", "erp", 22, "text"),
        ("ERP Status", "erp_status", 11, "text"), ("ERP Evidence", "erp_evidence", 40, "wrap"),
        ("Third-Party Apps", "third_party_apps_text", 44, "wrap"),
        ("Existing S2P Product", "existing_s2p_product", 18, "text"),
        ("S2P Product Detail", "existing_s2p_detail", 26, "wrap"),
        ("S2P Platform Status", "s2p_platform_status", 20, "text"),
        ("S2P Signal Level", "s2p_signal_level", 20, "text"),
        ("S2P Strong Signals", "s2p_strong_signals", 60, "wrap"),
        ("Digital Transformation Signals", "digital_transformation_signals", 44, "wrap"),
        ("Procurement Transformation Signals", "procurement_transformation_signals", 44, "wrap"),
        ("Relevant Technologies", "relevant_technologies", 30, "wrap"),
        ("Known Implementation Partner", "known_implementation_partner", 22, "wrap"),
        ("Known Consulting Partner", "known_consulting_partner", 22, "wrap"),
        ("Coupa Opportunity Type", "coupa_opportunity_type", 26, "text"),
        ("Ariba Opportunity Type", "ariba_opportunity_type", 26, "text"),
        ("Potential Opportunity", "potential_opportunity", 50, "wrap"),
        ("Account Notes", "account_notes", 50, "wrap"),
        ("Research Channel", "research_channel", 11, "text"),
        ("First Found", "first_found", 11, "text"), ("Last Researched", "last_researched", 11, "text"),
        ("Research Confidence", "research_confidence", 11, "text"),
        ("Account Owner", "account_owner", 16, "text"), ("Account Priority", "account_priority", 12, "text"),
        ("Pitch Angle / Next Step", "pitch_next_step", 36, "wrap"),
    ]
    write_table(ws, "tblAccounts", "ACCOUNTS — UAE Listed ICP (Revenue > USD 250M, 100+ employees)",
                f"Research Channel: Claude · Research date {RUN_DATE} · Gold columns are for your input",
                acc_cols, A, input_keys={"account_owner", "account_priority", "pitch_next_step"}, tab=1,
                validations={"account_priority": ["P1", "P2", "P3", "Park"]})

    # ---------------- Contacts ----------------
    ws = wb.create_sheet("Contacts")
    con_cols = [
        ("Contact ID", "contact_id", 11, "text"), ("Company", "company", 26, "text"),
        ("Full Name", "full_name", 22, "text"), ("Nationality", "nationality", 14, "text"),
        ("Title (Verbatim)", "title_verbatim", 32, "wrap"), ("Standardized Title", "standardized_title", 22, "text"),
        ("Role Family", "role_family", 14, "text"), ("Contact Tier", "contact_tier", 9, "text"),
        ("Research Channel", "research_channel", 10, "text"), ("Channel Source", "channel_source", 18, "text"),
        ("Source", "source", 22, "wrap"), ("Source Type", "source_type", 18, "text"),
        ("Source URL", "source_url", 26, "url"), ("Second Source URL", "second_source_url", 26, "url"),
        ("Verification Status", "verification_status", 13, "text"),
        ("Email", "email", 28, "text"), ("Email Status", "email_status", 14, "text"),
        ("Email Source", "email_source", 12, "text"), ("Email Confidence", "email_confidence", 14, "text"),
        ("Unverified Email Candidate", "email_candidate", 30, "wrap"),
        ("Phone", "phone", 17, "text"), ("Phone Type", "phone_type", 11, "text"),
        ("Phone Source", "phone_source", 12, "text"),
        ("Location", "location", 18, "text"), ("Country", "country", 10, "text"),
        ("LinkedIn URL", "linkedin_url", 30, "url"), ("Company Website", "company_website", 22, "url"),
        ("Existing S2P Product", "existing_s2p_product", 16, "text"),
        ("S2P Platform Status", "s2p_platform_status", 18, "text"),
        ("Account S2P Signal", "account_s2p_signal", 18, "text"),
        ("Contact S2P Signal", "s2p_contact_signal", 34, "wrap"),
        ("Coupa Opportunity Type", "coupa_opportunity_type", 24, "text"),
        ("Ariba Opportunity Type", "ariba_opportunity_type", 24, "text"),
        ("Notes / Intel About Contact", "notes_contact", 50, "wrap"),
        ("Research Date", "research_date", 11, "text"), ("Confidence", "confidence", 10, "text"),
        ("Record Status", "record_status", 18, "text"), ("Claude Check (vs reference)", "claude_check", 34, "wrap"),
        ("Owner", "owner", 12, "text"), ("Warm Intro?", "warm_intro", 10, "text"),
        ("Company (as in reference)", "company_ref_name", 22, "text"),
        ("Employment Status", "employment_status", 13, "text"),
        ("Previous Company", "previous_company", 18, "text"), ("Previous Title", "previous_title", 20, "text"),
        ("Campaign", "campaign", 16, "text"), ("Review Status", "review_status", 13, "text"),
        ("Reviewer Notes", "reviewer_notes", 30, "wrap"),
    ]
    tiers = ["Tier 1", "Tier 2", "Tier 3", "Tier 4"]
    families = ["PROCUREMENT", "FINANCE", "IT", "SUPPLY CHAIN", "TRANSFORMATION", "OTHER"]
    write_table(ws, "tblContacts", "CONTACTS — Decision Makers & Influencers",
                "Channel State = Claude for every row Claude added · Emails never pattern-guessed · Gold columns are for your input",
                con_cols, P, input_keys={"campaign", "review_status", "reviewer_notes"}, tab=2,
                validations={"contact_tier": tiers, "role_family": families,
                             "review_status": ["To Review", "Approved", "Rejected", "Needs Research"]})

    # ---------------- Stakeholders (campaign view) ----------------
    ws = wb.create_sheet("Stakeholders")
    stk_cols = [
        ("Company", "company", 26, "text"), ("Full Name", "full_name", 22, "text"),
        ("Nationality", "nationality", 14, "text"), ("Title (Verbatim)", "title_verbatim", 32, "wrap"),
        ("Role Family", "role_family", 14, "text"), ("Contact Tier", "contact_tier", 9, "text"),
        ("Channel State", "channel_state", 10, "text"), ("Channel Source", "channel_source", 18, "text"),
        ("Email", "email", 28, "text"), ("Email Status", "email_status", 14, "text"),
        ("Unverified Email Candidate", "email_candidate", 30, "wrap"),
        ("Phone", "phone", 17, "text"), ("Phone Type", "phone_type", 11, "text"),
        ("Location", "location", 18, "text"), ("LinkedIn URL", "linkedin_url", 30, "url"),
        ("Existing S2P Product", "existing_s2p_product", 16, "text"),
        ("S2P Signal", "account_s2p_signal", 18, "text"),
        ("Employment Status", "employment_status", 13, "text"),
        ("Verification Status", "verification_status", 13, "text"),
        ("Notes / Intel about the Contact", "notes_contact", 48, "wrap"),
        ("Notes / Intel about the Company", "notes_company", 60, "wrap"),
        ("Record Status", "record_status", 18, "text"), ("Claude Check (vs reference)", "claude_check", 34, "wrap"),
        ("Owner", "owner", 12, "text"), ("Warm Intro?", "warm_intro", 10, "text"),
        ("Campaign", "campaign", 16, "text"), ("Outreach Status", "outreach_status", 14, "text"),
    ]
    tier_rank = {t: i for i, t in enumerate(tiers)}
    sig_rank = {s: i for i, s in enumerate(SIGNAL_ORDER)}
    stk_rows = sorted(P, key=lambda p: (sig_rank.get(p.get("account_s2p_signal"), 9), p["company"],
                                        tier_rank.get(p.get("contact_tier"), 9), name_key(p.get("full_name")),
                                        0 if p.get("_is_ref") else 1))
    write_table(ws, "tblStakeholders", "STAKEHOLDERS — Campaign Planning View",
                "Company & contact details only · sorted by account S2P signal strength, then tier · Gold columns are for your input",
                stk_cols, stk_rows, input_keys={"campaign", "outreach_status", "owner", "warm_intro"}, tab=3,
                validations={"outreach_status": ["Not Started", "Queued", "Contacted", "Replied", "Meeting Booked", "Not Interested"],
                             "contact_tier": tiers})

    # ---------------- S2P Signals ----------------
    ws = wb.create_sheet("S2P Signals")
    sig_cols = [("Company ID", "company_id", 10, "text"), ("Company", "company", 26, "text"),
                ("Category", "category", 20, "text"), ("Signal", "signal", 50, "wrap"),
                ("Level", "level", 20, "text"), ("Platform", "platform", 12, "text"),
                ("Evidence", "evidence", 60, "wrap"), ("Source URL", "source_url", 30, "url"),
                ("Date", "date", 11, "text"), ("Research Channel", "research_channel", 10, "text")]
    sigs = sorted(d["signals"], key=lambda s: (sig_rank.get(s.get("level"), 9), s["company"]))
    write_table(ws, "tblSignals", "S2P SIGNALS — Evidence-based, scoring-free",
                "Every signal carries its evidence and source · VERY STRONG → WEAK", sig_cols, sigs, tab=4)

    # ---------------- ERP & Apps ----------------
    ws = wb.create_sheet("ERP & Apps Landscape")
    erp_rows = [{"company_id": a["company_id"], "company": a["company_name"], "name": a.get("erp", ""),
                 "category": "ERP (core)", "status": a.get("erp_status", ""), "evidence": a.get("erp_evidence", ""),
                 "source_url": ""} for a in A]
    app_rows = erp_rows + d["apps"]
    app_rows.sort(key=lambda r: (r["company"], r.get("category") != "ERP (core)"))
    app_cols = [("Company ID", "company_id", 10, "text"), ("Company", "company", 26, "text"),
                ("Application / Platform", "name", 26, "text"), ("Category", "category", 18, "text"),
                ("Status", "status", 12, "text"), ("Evidence", "evidence", 60, "wrap"),
                ("Source URL", "source_url", 30, "url")]
    write_table(ws, "tblApps", "ERP LANDSCAPE & THIRD-PARTY APPLICATIONS",
                "FACT = directly sourced · LIKELY = multiple indirect signals · UNVERIFIED = single weak source (e.g. technographics)",
                app_cols, app_rows, tab=5)

    # ---------------- Source Evidence ----------------
    ws = wb.create_sheet("Source Evidence")
    src_cols = [("Company ID", "company_id", 10, "text"), ("Company", "company", 24, "text"),
                ("Related Contact", "related_contact", 20, "text"), ("Source", "source", 24, "wrap"),
                ("Source Type", "source_type", 18, "text"), ("Source Tier", "source_tier", 9, "text"),
                ("Source URL", "url", 30, "url"), ("Research Channel", "research_channel", 10, "text"),
                ("Information Found", "information_found", 40, "wrap"), ("Evidence", "evidence", 50, "wrap"),
                ("Date Published", "date_published", 11, "text"), ("Date Accessed", "date_accessed", 11, "text"),
                ("Confidence", "confidence", 10, "text"),
                ("Supports Current Employment?", "supports_current_employment", 11, "text"),
                ("Supports Current Title?", "supports_current_title", 11, "text"),
                ("Supports S2P Status?", "supports_s2p_status", 11, "text")]
    write_table(ws, "tblSources", "SOURCE EVIDENCE — Audit Trail",
                "How do we know this? Every observation retained, never deleted", src_cols, d["sources"], tab=6)

    # ---------------- Employment History ----------------
    ws = wb.create_sheet("Employment History")
    eh_cols = [("Company ID", "company_id", 10, "text"), ("Full Name", "full_name", 22, "text"),
               ("Company", "company", 26, "text"), ("Title", "title", 30, "wrap"),
               ("Source", "source", 22, "text"), ("Source URL", "source_url", 30, "url"),
               ("Date Found", "date_found", 11, "text"), ("Determination", "determination", 26, "text"),
               ("Evidence", "evidence", 50, "wrap")]
    write_table(ws, "tblHistory", "EMPLOYMENT HISTORY", "Current vs previous vs recently changed — with evidence",
                eh_cols, d["history"], tab=7)

    # ---------------- Conflicts ----------------
    ws = wb.create_sheet("Conflicts")
    cf_cols = [("Company ID", "company_id", 10, "text"), ("Company", "company", 24, "text"),
               ("Entity", "entity", 22, "text"), ("Field", "field", 12, "text"),
               ("Value A", "value_a", 28, "wrap"), ("Source A", "source_a", 22, "wrap"),
               ("Value B", "value_b", 28, "wrap"), ("Source B", "source_b", 22, "wrap"),
               ("Determination", "determination", 30, "wrap"), ("Evidence", "evidence", 40, "wrap"),
               ("Resolution (your input)", "resolution", 26, "wrap")]
    write_table(ws, "tblConflicts", "CONFLICTS — Both values retained", "Resolve in the gold column; nothing is overwritten",
                cf_cols, d["conflicts"], input_keys={"resolution"}, tab=8)

    # ---------------- Research Summary ----------------
    ws = wb.create_sheet("Research Summary")
    summ = []
    for a in A:
        cs = by_co.get(a["company_id"], [])
        summ.append({
            "company": a["company_name"], "country": a.get("country"), "industry": a.get("industry"),
            "revenue": a.get("revenue_usd_m"), "employees": a.get("employee_range"),
            "s2p": a.get("existing_s2p_product"), "status": a.get("s2p_platform_status"),
            "sig": a.get("s2p_signal_level"), "signals": a.get("s2p_strong_signals"),
            "contacts": "\n".join(f"{p.get('full_name')} | {p.get('title_verbatim')} | {p.get('role_family')} | "
                                  f"{p.get('contact_tier')} | {p.get('email_status')}" for p in cs),
            "opp": a.get("potential_opportunity"),
            "evidence": "\n".join(f"{s.get('source')} ({s.get('date_published','')}) — {s.get('information_found','')[:90]}"
                                  for s in d["sources"] if s["company_id"] == a["company_id"])[:1500],
            "conf": a.get("research_confidence")})
    sm_cols = [("Company", "company", 26, "text"), ("Country", "country", 8, "text"),
               ("Industry", "industry", 18, "text"), ("Revenue (USD m)", "revenue", 11, "num"),
               ("Employees", "employees", 12, "text"), ("Current S2P", "s2p", 16, "text"),
               ("S2P Status", "status", 18, "text"), ("Signal", "sig", 18, "text"),
               ("Strong Signals", "signals", 60, "wrap"), ("Key Contacts", "contacts", 60, "wrap"),
               ("Opportunity Observations", "opp", 50, "wrap"), ("Evidence", "evidence", 60, "wrap"),
               ("Confidence", "conf", 10, "text")]
    write_table(ws, "tblSummary", "RESEARCH SUMMARY — One row per account", "", sm_cols, summ, tab=9)
    COLMAP["Research Summary"]["s2p_signal_level"] = COLMAP["Research Summary"]["sig"]

    # ---------------- Settings ----------------
    ws = wb.create_sheet("Settings")
    ws.sheet_properties.tabColor = TAB_COLORS[12]
    banner(ws, "SETTINGS — Configurable Contact Tiering & Legend", "Edit the gold framed cells; tiers are an internal sales classification, not a fact", 4)
    rows = [("Tier", "Label", "Example titles (editable)"),
            ("Tier 1", "Primary Decision Maker", "CPO, CFO, CIO, CTO, CDO, Chief Transformation Officer"),
            ("Tier 2", "Executive Influencer", "VP Procurement, Procurement Director, Finance Director, IT Director, Chief Supply Chain Officer"),
            ("Tier 3", "Functional Influencer", "Head of Procurement, Sourcing Director, ERP Director, Procurement Transformation Lead, Head of Shared Services"),
            ("Tier 4", "Operational / Technical", "Procurement Manager, S2P Manager, Coupa/Ariba Administrator, Procurement Systems Manager")]
    for ri, r in enumerate(rows, 4):
        for ci, v in enumerate(r, 1):
            c = ws.cell(row=ri, column=ci, value=v)
            c.font = Font(name=FONT, size=10, bold=ri == 4, color="FFFFFF" if ri == 4 else "1B263B")
            c.alignment = Alignment(wrap_text=True, vertical="top")
            if ri == 4:
                c.fill = PatternFill("solid", fgColor=HEADER)
            elif ci > 1:
                c.fill = PatternFill("solid", fgColor=INPUT_FILL)
                c.border = INPUT_BORDER_STYLE
    legend = [("LEGEND", ""), ("Gold framed cells / gold headers", "Your input — safe to edit"),
              ("Channel State = Claude", "Row added by Claude research"),
              ("Channel State = Copilot", "Row supplied by Copilot (never modified)"),
              ("Email Status 'Verified Active'", "Data-provider (Seamless.ai) verification flag"),
              ("Email Status 'Publicly Listed'", "Email printed on an official/reputable page for that person"),
              ("Blank email", "Could not be reliably verified — deliberately left blank"),
              ("Nationality 'Not Publicly Verified'", "Never inferred from name, photo or location"),
              ("FACT / LIKELY / UNVERIFIED / CONFLICTING / UNKNOWN", "Verification vocabulary used throughout")]
    for ri, (k, v) in enumerate(legend, 11):
        ws.cell(row=ri, column=1, value=k).font = Font(name=FONT, size=10, bold=ri == 11)
        ws.cell(row=ri, column=2, value=v).font = Font(name=FONT, size=10)
    ws.column_dimensions["A"].width = 44
    ws.column_dimensions["B"].width = 44
    ws.column_dimensions["C"].width = 80

    # ---------------- Pivot Analysis ----------------
    ws = wb.create_sheet("Pivot Analysis")
    ws.sheet_properties.tabColor = TAB_COLORS[10]
    banner(ws, "PIVOT ANALYSIS — Live summary tables", "All counts are formulas over the Accounts / Contacts / History tabs", 14)
    ACC, CON = "Accounts", "Contacts"
    pivots = []

    def pivot(r0, c0, title, labels, formula_fn, header=("Value", "Count")):
        ws.cell(row=r0, column=c0, value=title).font = Font(name=FONT, bold=True, size=11, color=HEADER)
        for i, h in enumerate(header):
            c = ws.cell(row=r0 + 1, column=c0 + i, value=h)
            c.fill = PatternFill("solid", fgColor=HEADER_2)
            c.font = Font(name=FONT, bold=True, color="FFFFFF", size=9)
        for i, lab in enumerate(labels):
            ws.cell(row=r0 + 2 + i, column=c0, value=lab).font = Font(name=FONT, size=9)
            c = ws.cell(row=r0 + 2 + i, column=c0 + 1, value=formula_fn(lab))
            c.font = Font(name=FONT, size=9)
            for j in range(2):
                ws.cell(row=r0 + 2 + i, column=c0 + j).border = BORDER
        pivots.append((title, r0, c0, len(labels)))
        return r0 + len(labels) + 4

    def uniq(rows, key, default="Unknown"):
        return sorted({(r.get(key) or default) for r in rows})

    L = lambda s, k: col(s, k)  # noqa: E731
    r = 4
    r = pivot(r, 1, "1. Country × Accounts", uniq(A, "country"), lambda v: f'=COUNTIF({L(ACC,"country")},"{v}")')
    r = pivot(r, 1, "2. S2P Platform × Accounts", uniq(A, "existing_s2p_product"),
              lambda v: f'=COUNTIF({L(ACC,"existing_s2p_product")},"{v}")')
    r = pivot(r, 1, "4. Role Family × Contacts", families, lambda v: f'=COUNTIF({L(CON,"role_family")},"{v}")')
    r = pivot(r, 1, "5. S2P Status × Accounts", uniq(A, "s2p_platform_status"),
              lambda v: f'=COUNTIF({L(ACC,"s2p_platform_status")},"{v}")')
    r = pivot(r, 1, "6. Coupa vs Ariba vs Other/Unknown", ["*Coupa*", "*Ariba*", "Other S2P", "No Evidence", "Unknown"],
              lambda v: f'=COUNTIF({L(ACC,"existing_s2p_product")},"{v}")')
    r2 = 4
    r2 = pivot(r2, 5, "7. Signal Level × Accounts", SIGNAL_ORDER,
               lambda v: f'=COUNTIF({L(ACC,"s2p_signal_level")},"{v}")')
    r2 = pivot(r2, 5, "8. Research Channel × Contacts", ["Claude", "Copilot", "Claude (prior run)", "Reference"],
               lambda v: f'=COUNTIF({L(CON,"research_channel")},"{v}")')
    r2 = pivot(r2, 5, "Contact Tier × Contacts", tiers, lambda v: f'=COUNTIF({L(CON,"contact_tier")},"{v}")')
    r2 = pivot(r2, 5, "Email Status × Contacts", ["Verified Active", "Publicly Listed", "Unverified", "Inactive", "Not Found"],
               lambda v: f'=COUNTIF({L(CON,"email_status")},"{v}")')
    r2 = pivot(r2, 5, "9. Employment Changes", ["*change*", "Previous employment", "Current employment"],
               lambda v: f'=COUNTIF(\'Employment History\'!$H${FIRST_ROW}:$H${MAX_ROW},"{v}")')
    r2 = pivot(r2, 5, "10a. Coupa Opportunity Type", uniq(A, "coupa_opportunity_type", "No Evidence"),
               lambda v: f'=COUNTIF({L(ACC,"coupa_opportunity_type")},"{v}")')
    r2 = pivot(r2, 5, "10b. Ariba Opportunity Type", uniq(A, "ariba_opportunity_type", "No Evidence"),
               lambda v: f'=COUNTIF({L(ACC,"ariba_opportunity_type")},"{v}")')
    r3 = 4
    r3 = pivot(r3, 9, "3. Company × Contacts", [a["company_name"] for a in sorted(A, key=lambda a: a["company_name"])],
               lambda v: f'=COUNTIF({L(CON,"company")},"{v.replace(chr(34), "")}")')
    r3 = pivot(r3, 9, "Industry × Accounts", uniq(A, "industry"), lambda v: f'=COUNTIF({L(ACC,"industry")},"{v}")')
    for cL, w in (("A", 34), ("B", 9), ("E", 34), ("F", 9), ("I", 38), ("J", 9)):
        ws.column_dimensions[cL].width = w

    # charts
    def add_bar(title, r0, c0, n, anchor, horizontal=True):
        ch = BarChart()
        ch.type = "bar" if horizontal else "col"
        ch.title = title
        ch.style = 10
        ch.legend = None
        data = Reference(ws, min_col=c0 + 1, min_row=r0 + 1, max_row=r0 + 1 + n)
        cats = Reference(ws, min_col=c0, min_row=r0 + 2, max_row=r0 + 1 + n)
        ch.add_data(data, titles_from_data=True)
        ch.set_categories(cats)
        ch.series[0].graphicalProperties.solidFill = "2E75B6"
        ch.dataLabels = DataLabelList()
        ch.dataLabels.showVal = True
        ch.height, ch.width = 7.5, 14
        return ch, anchor

    charts = {p[0]: p for p in pivots}
    placements = [("7. Signal Level × Accounts", "L4"), ("2. S2P Platform × Accounts", "L20"),
                  ("4. Role Family × Contacts", "L36"), ("Industry × Accounts", "L52")]
    for title, anchor in placements:
        t, r0, c0, n = charts[title]
        ch, a = add_bar(title, r0, c0, n, anchor)
        ws.add_chart(ch, a)

    # ---------------- Executive Dashboard ----------------
    ws = ws_dash
    ws.sheet_properties.tabColor = TAB_COLORS[0]
    banner(ws, "ACCOUNT COPILOT — B2B Procurement Intelligence · UAE",
           f"ICP: UAE stock-listed (ADX/DFM) · Net revenue > USD 250M · 100+ employees · Research date {RUN_DATE} · Research Channel: Claude", 12)
    for cL in "ABCDEFGHIJKL":
        ws.column_dimensions[cL].width = 16
    kpis = [
        ("Companies researched", f'=COUNTA({L(ACC,"company_name")})'),
        ("ICP fit = Yes", f'=COUNTIF({L(ACC,"icp_fit")},"Yes")'),
        ("Contacts identified", f'=COUNTA({L(CON,"full_name")})'),
        ("Verified contacts", f'=COUNTIF({L(CON,"verification_status")},"VERIFIED")'),
        ("Contacts requiring verification", f'=COUNTA({L(CON,"full_name")})-COUNTIF({L(CON,"verification_status")},"VERIFIED")'),
        ("Contacts with email", f'=COUNTIF({L(CON,"email")},"?*")'),
        ("Existing Coupa accounts", f'=COUNTIF({L(ACC,"existing_s2p_product")},"*Coupa*")'),
        ("Existing Ariba accounts", f'=COUNTIF({L(ACC,"existing_s2p_product")},"*Ariba*")'),
        ("S2P transformation signals", f'=COUNTA(\'S2P Signals\'!$D${FIRST_ROW}:$D${MAX_ROW})'),
        ("Strong / very strong accounts", f'=COUNTIF({L(ACC,"s2p_signal_level")},"STRONG SIGNAL")+COUNTIF({L(ACC,"s2p_signal_level")},"VERY STRONG SIGNAL")'),
        ("Possible new S2P projects", "=" + "+".join(f'COUNTIF({L(ACC,"s2p_platform_status")},"{s}")' for s in
                                                    ["Evaluation", "RFP / Tender", "Currently Implementing", "Replacement / Transformation"])),
        ("Managed-services opportunities", f'=COUNTIF({L(ACC,"coupa_opportunity_type")},"*Managed*")+COUNTIF({L(ACC,"ariba_opportunity_type")},"*Managed*")'),
        ("Employment changes detected", f'=COUNTIF(\'Employment History\'!$H${FIRST_ROW}:$H${MAX_ROW},"*change*")'),
        ("Conflicting records", f'=COUNTA(Conflicts!$D${FIRST_ROW}:$D${MAX_ROW})'),
        ("Source observations", f'=COUNTA(\'Source Evidence\'!$G${FIRST_ROW}:$G${MAX_ROW})'),
        ("Emails verified active", f'=COUNTIF({L(CON,"email_status")},"Verified Active")'),
    ]
    tile_fill = PatternFill("solid", fgColor="0F2744")
    for i, (label, f) in enumerate(kpis):
        rr = 4 + (i // 4) * 4
        cc = 1 + (i % 4) * 3
        ws.merge_cells(start_row=rr, start_column=cc, end_row=rr, end_column=cc + 2)
        ws.merge_cells(start_row=rr + 1, start_column=cc, end_row=rr + 2, end_column=cc + 2)
        lc = ws.cell(row=rr, column=cc, value=label.upper())
        vc = ws.cell(row=rr + 1, column=cc, value=f)
        for r_ in range(rr, rr + 3):
            for c_ in range(cc, cc + 3):
                ws.cell(row=r_, column=c_).fill = tile_fill
        lc.font = Font(name=FONT, size=8, bold=True, color="8DB9EC")
        vc.font = Font(name=FONT, size=24, bold=True, color="FFFFFF")
        lc.alignment = Alignment(horizontal="left", indent=1, vertical="bottom")
        vc.alignment = Alignment(horizontal="left", indent=1, vertical="center")
    # top accounts list
    top = [a for a in A if a.get("s2p_signal_level") in ("VERY STRONG SIGNAL", "STRONG SIGNAL")]
    top.sort(key=lambda a: sig_rank.get(a.get("s2p_signal_level"), 9))
    r0 = 21
    ws.cell(row=r0, column=1, value="PRIORITY ACCOUNTS — Strong S2P signals").font = Font(name=FONT, bold=True, size=12, color=HEADER)
    hdr = ["Company", "Existing S2P", "S2P Status", "Signal", "ERP", "Why (evidence summary)"]
    spans = [2, 2, 2, 2, 1, 3]
    cc = 1
    for h, s in zip(hdr, spans):
        ws.merge_cells(start_row=r0 + 1, start_column=cc, end_row=r0 + 1, end_column=cc + s - 1)
        c = ws.cell(row=r0 + 1, column=cc, value=h)
        c.fill = PatternFill("solid", fgColor=HEADER)
        c.font = Font(name=FONT, bold=True, color="FFFFFF", size=9)
        cc += s
    for i, a in enumerate(top[:25]):
        rr = r0 + 2 + i
        vals = [a["company_name"], a.get("existing_s2p_product"), a.get("s2p_platform_status"),
                a.get("s2p_signal_level"), a.get("erp"), (a.get("s2p_strong_signals") or "")[:220]]
        cc = 1
        for v, s in zip(vals, spans):
            ws.merge_cells(start_row=rr, start_column=cc, end_row=rr, end_column=cc + s - 1)
            c = ws.cell(row=rr, column=cc, value=v)
            c.font = Font(name=FONT, size=9)
            c.alignment = Alignment(wrap_text=True, vertical="top")
            c.fill = PatternFill("solid", fgColor=BAND if i % 2 == 0 else "FFFFFF")
            if v in SIGNAL_FILLS:
                bg, fg = SIGNAL_FILLS[v]
                c.fill = PatternFill("solid", fgColor=bg)
                c.font = Font(name=FONT, size=9, bold=True, color=fg)
            cc += s
        ws.row_dimensions[rr].height = 48
    t, pr, pc, n = charts["7. Signal Level × Accounts"]
    pie = PieChart()
    pie.title = "Accounts by S2P signal"
    pie.add_data(Reference(wb["Pivot Analysis"], min_col=pc + 1, min_row=pr + 1, max_row=pr + 1 + n), titles_from_data=True)
    pie.set_categories(Reference(wb["Pivot Analysis"], min_col=pc, min_row=pr + 2, max_row=pr + 1 + n))
    pie.height, pie.width = 8, 12
    ws.add_chart(pie, "N4")
    ws.sheet_view.showGridLines = False

    order = ["Executive Dashboard", "Accounts", "Contacts", "Stakeholders", "S2P Signals", "ERP & Apps Landscape",
             "Source Evidence", "Employment History", "Conflicts", "Research Summary", "Pivot Analysis", "Settings"]
    wb._sheets = [wb[n] for n in order]
    for s in wb.worksheets:
        s.sheet_view.zoomScale = 90
    from openpyxl.workbook.properties import CalcProperties
    wb.calculation = CalcProperties(fullCalcOnLoad=True)
    wb.save(out)
    print(f"Saved {out}: {len(A)} accounts, {len(P)} contacts, {len(d['signals'])} signals, "
          f"{len(d['sources'])} sources, {len(d['conflicts'])} conflicts")


if __name__ == "__main__":
    out = sys.argv[1] if len(sys.argv) > 1 else os.path.join(ROOT, "output", "Account_CoPilot_Master_Book_UAE.xlsx")
    main(out)
