"""Import the reference Stakeholders worksheet (Copilot + earlier Claude rows) without changing any value.

Usage: python3 scripts/import_reference.py "<path to FINAL-UAE-Target-LIST.xlsx>"
Writes data/reference/stakeholders_ref.json (rows mapped to researched accounts) and
data/reference/unmapped_companies.txt (companies outside the researched ICP account list).
"""
import glob
import json
import os
import re
import sys

from openpyxl import load_workbook

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from load_data import ROOT, RESEARCH_DIR  # noqa: E402

STOP = {"pjsc", "psc", "plc", "llc", "pjs", "p", "j", "s", "c", "co", "company", "the", "holding", "holdings",
        "group", "ltd", "limited", "formerly", "parent", "incl", "corporation", "of", "and", "bank"}
# reference-name -> slug where names are too different to match automatically
MANUAL = {
    "e&": "emirates-telecommunications-group", "e& uae": "emirates-telecommunications-group",
    "e& (etisalat group)": "emirates-telecommunications-group",
    "taqa group": "taqa", "abu dhabi national energy company pjsc": "taqa",
    "mashreq bank": "mashreq", "mashreqbank psc": "mashreq",
    "abu dhabi commercial bank": "adcb", "park in": "parkin-company",
    "adnoc drilling": "adnoc-drilling", "adnh catering": "adnh-catering",
    "gulf pharmaceutical industries psc": "julphar", "lulu retail": "lulu-retail-holdings",
    "spinneys": "spinneys-1961-holding", "taaleem": "taaleem-holdings",
    "al ansari exchange": "al-ansari-financial-services",
    "emirates steel arkan (emsteel)": "emsteel", "tabreed": "national-central-cooling-company-tabreed",
    "national central cooling pjsc (tabreed)": "national-central-cooling-company-tabreed",
    "dewa (dubai electricity & water)": "dubai-electricity-and-water-authority",
    "purehealth": "pure-health", "abu dhabi aviation (ada)": "abu-dhabi-aviation",
    "alec holdings (formerly alec construction)": "alec-holdings",
}
NOT_SAME = {"al-futtaim group", "dnata", "six construct llc (besix)", "dubai holding (incl. nakheel meydan meraas dubai properties)", "dubai holdings", "emirates group", "bank of sharjah pjsc", "اتحاد الإمارات للرياضات البحرية", "e& international", "moro hub (digital dewa)", "moro hub", "emaar hospitality group", "adnoc (parent)"}


def toks(s):
    s = re.sub(r"\(.*?\)", " ", s.lower())
    return [t for t in re.findall(r"[a-z0-9&]+", s) if t not in STOP]


def main(path):
    accounts = {}
    for f in glob.glob(os.path.join(RESEARCH_DIR, "*.json")):
        slug = os.path.basename(f)[:-5]
        c = json.load(open(f)).get("company", {})
        name = c.get("company_name", slug)
        abbrevs = [a.lower() for a in re.findall(r"\(([^)]+)\)", name)]
        accounts[slug] = {"name": name, "key": " ".join(toks(name)), "abbr": abbrevs,
                          "domain": (c.get("domain") or "").split(" ")[0].lower()}
    wb = load_workbook(path, read_only=True)
    ws = wb["Stakeholders"]
    header = None
    out, unmapped = [], {}
    for i, row in enumerate(ws.iter_rows(values_only=True), 1):
        if header is None:
            if row and "Full name" in [str(x) for x in row]:
                header = [str(x or "").strip() for x in row]
            continue
        rec = dict(zip(header, row))
        if not rec.get("Full name") or not rec.get("Company") or str(rec.get("Company")).startswith("Company ("):
            continue
        if isinstance(rec["Full name"], (int, float)) or str(rec["Full name"]).strip().isdigit():
            continue  # annotation row, not a contact
        co = str(rec["Company"]).strip()
        lc = co.lower()
        slug = None
        if lc in NOT_SAME:
            slug = None
        elif lc in MANUAL:
            slug = MANUAL[lc] if MANUAL[lc] in accounts else None
        else:
            k = " ".join(toks(co))
            abbr = [a.lower() for a in re.findall(r"\(([^)]+)\)", co)]
            em = str(rec.get("Email") or "").lower()
            dom = em.split("@")[1] if "@" in em else ""
            for s, a in accounts.items():
                kt, at = k.split(), a["key"].split()
                if k and (k == a["key"] or (min(len(kt), len(at)) >= 2 and (at[:len(kt)] == kt or kt[:len(at)] == at))):
                    slug = s
                    break
                if abbr and set(abbr) & set(a["abbr"]):
                    slug = s
                    break
            if not slug and dom:
                for s, a in accounts.items():
                    if a["domain"] and dom == a["domain"]:
                        slug = s
                        break
        rec = {k: ("" if v is None else v) for k, v in rec.items()}
        rec["_ref_row"] = i
        rec["_slug"] = slug
        if slug:
            out.append(rec)
        else:
            unmapped[co] = unmapped.get(co, 0) + 1
    os.makedirs(os.path.join(ROOT, "data", "reference"), exist_ok=True)
    json.dump(out, open(os.path.join(ROOT, "data", "reference", "stakeholders_ref.json"), "w"), indent=1, default=str)
    with open(os.path.join(ROOT, "data", "reference", "unmapped_companies.txt"), "w") as f:
        for k, v in sorted(unmapped.items()):
            f.write(f"{v}\t{k}\n")
    mapped = {}
    for r in out:
        mapped.setdefault(r["_slug"], set()).add(r["Company"])
    for s, names in sorted(mapped.items()):
        print(f"{s:45s} <- {sorted(names)}")
    print(f"\nMapped rows: {len(out)} · unmapped companies: {len(unmapped)} ({sum(unmapped.values())} rows)")


if __name__ == "__main__":
    main(sys.argv[1])
